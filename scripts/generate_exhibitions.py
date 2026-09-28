from __future__ import annotations

from datetime import date, datetime
from html import escape
from pathlib import Path
import re
import sys
from urllib.parse import urlparse

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
WORKBOOK = ROOT / "Bfoxyarts_Exhibitions_Content_Template.xlsx"
OUTPUT = ROOT / "exhibitions.html"


def text(value) -> str:
    return "" if value is None else str(value).strip()


def yes(value) -> bool:
    return text(value).casefold() in {"yes", "y", "true", "1"}


def valid_url(value) -> str:
    value = text(value)
    if not value:
        return ""
    parsed = urlparse(value)
    return value if parsed.scheme in {"http", "https"} and parsed.netloc else ""


def excel_date(value, field: str, row_number: int) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        for fmt in ("%Y-%m-%d", "%m/%d/%Y", "%m/%d/%y"):
            try:
                return datetime.strptime(value.strip(), fmt).date()
            except ValueError:
                pass
    raise ValueError(f"Row {row_number}: {field} must contain a real Excel date.")


def slugify(value: str) -> str:
    slug = re.sub(r"[^a-z0-9]+", "-", value.casefold()).strip("-")
    return slug or "exhibition"


def display_date(start: date, end: date) -> str:
    if start == end:
        return f"{start.strftime('%B')} {start.day}, {start.year}"
    if start.year == end.year and start.month == end.month:
        return f"{start.strftime('%B')} {start.day}–{end.day}, {start.year}"
    if start.year == end.year:
        return f"{start.strftime('%B')} {start.day}–{end.strftime('%B')} {end.day}, {start.year}"
    return f"{start.strftime('%B')} {start.day}, {start.year}–{end.strftime('%B')} {end.day}, {end.year}"


def paragraphs(value: str) -> str:
    blocks = [x.strip() for x in re.split(r"\n\s*\n|\r\n\s*\r\n", value) if x.strip()]
    if not blocks and value.strip():
        blocks = [value.strip()]
    return "".join(f'<p class="exhibition-description">{escape(block).replace(chr(10), "<br>")}</p>' for block in blocks)


def image_markup(row: dict[str, object]) -> str:
    images = []
    for number in (1, 2):
        filename = text(row.get(f"Image Filename {number}"))
        if not filename:
            continue
        alt = text(row.get(f"Image Alt Text {number}")) or f"Exhibition image for {text(row.get('Exhibition Name'))}"
        images.append(f'<img src="images/{escape(filename, quote=True)}" alt="{escape(alt, quote=True)}" loading="lazy">')
    if not images:
        return '<div class="exhibition-image-placeholder" aria-hidden="true">Bfoxyarts</div>'
    css = "exhibition-media contain" + (" duo" if len(images) > 1 else "")
    return f'<div class="{css}">{"".join(images)}</div>'


def action_links(row: dict[str, object]) -> str:
    links = []
    official = valid_url(row.get("Official Show / Advertisement URL"))
    venue = valid_url(row.get("Venue Website"))
    ticket = valid_url(row.get("Ticket / Registration URL"))
    if official:
        links.append(f'<a class="button" href="{escape(official, quote=True)}" target="_blank" rel="noopener noreferrer">Official show details</a>')
    if venue:
        links.append(f'<a class="text-link" href="{escape(venue, quote=True)}" target="_blank" rel="noopener noreferrer">Visit venue website</a>')
    if ticket and ticket not in {official, venue}:
        links.append(f'<a class="text-link" href="{escape(ticket, quote=True)}" target="_blank" rel="noopener noreferrer">Tickets or registration</a>')
    return f'<div class="exhibition-actions">{"".join(links)}</div>' if links else ""


def card_markup(row: dict[str, object], row_number: int) -> str:
    name = text(row.get("Exhibition Name"))
    venue = text(row.get("Venue / Gallery"))
    if not name or not venue:
        raise ValueError(f"Row {row_number}: published rows require Exhibition Name and Venue / Gallery.")
    start = excel_date(row.get("Start Date"), "Start Date", row_number)
    end = excel_date(row.get("End Date"), "End Date", row_number)
    if end < start:
        raise ValueError(f"Row {row_number}: End Date cannot be before Start Date.")
    street = text(row.get("Street Address"))
    city = text(row.get("City"))
    state = text(row.get("State"))
    zip_code = text(row.get("ZIP"))
    locality = ", ".join(x for x in (city, state) if x)
    address = ", ".join(x for x in (street, locality, zip_code) if x)
    reception_date = row.get("Opening Reception Date")
    reception = ""
    if isinstance(reception_date, (date, datetime)):
        rd = reception_date.date() if isinstance(reception_date, datetime) else reception_date
        reception = f"Opening reception: {rd.strftime('%B')} {rd.day}, {rd.year}"
    elif text(reception_date):
        reception = f"Opening reception: {text(reception_date)}"
    reception_time = text(row.get("Reception Time"))
    if reception and reception_time:
        reception += f", {reception_time}"
    artwork = text(row.get("Participating Artwork"))
    description = text(row.get("Short Description"))
    featured = " featured" if yes(row.get("Featured?")) else ""
    slug = text(row.get("URL Slug")) or slugify(name)
    meta = [f"<p><strong>{escape(venue)}</strong></p>"]
    if address:
        meta.append(f"<p>{escape(address)}</p>")
    meta.append(f"<p><strong>{escape(display_date(start, end))}</strong></p>")
    if reception:
        meta.append(f"<p>{escape(reception)}</p>")
    artwork_note = f'<p class="artwork-note"><strong>Bernadette\'s participating artwork:</strong> <em>{escape(artwork)}</em></p>' if artwork else ""
    return f'''<article class="exhibition-card{featured}" id="{escape(slug, quote=True)}" data-start="{start.isoformat()}" data-end="{end.isoformat()}">
{image_markup(row)}
<div class="exhibition-content">
<div class="status-row"><span class="status-badge"></span>{'<span class="featured-badge">Featured</span>' if featured else ''}</div>
<h2>{escape(name)}</h2>
<div class="exhibition-meta">{"".join(meta)}</div>
{paragraphs(description)}
{artwork_note}
{action_links(row)}
</div>
</article>'''


def load_settings(workbook) -> dict[str, str]:
    settings = {}
    if "Page Settings" not in workbook.sheetnames:
        return settings
    ws = workbook["Page Settings"]
    for row in ws.iter_rows(values_only=True):
        if len(row) >= 2 and text(row[0]):
            settings[text(row[0])] = text(row[1])
    return settings


def load_cards(workbook) -> list[str]:
    if "Exhibitions" not in workbook.sheetnames:
        raise ValueError("Workbook must contain an Exhibitions worksheet.")
    ws = workbook["Exhibitions"]
    headers = [text(cell.value) for cell in ws[1]]
    required = {"Publish?", "Featured?", "Exhibition Name", "Venue / Gallery", "Start Date", "End Date"}
    missing = sorted(required - set(headers))
    if missing:
        raise ValueError("Missing required columns: " + ", ".join(missing))
    cards = []
    for row_number, values in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        row = dict(zip(headers, values))
        if not yes(row.get("Publish?")):
            continue
        cards.append(card_markup(row, row_number))
    if not cards:
        raise ValueError("No rows have Publish? set to Yes.")
    return cards


def build_html(cards: list[str], settings: dict[str, str]) -> str:
    title = settings.get("Page Title") or "Exhibitions | Bernadette Fox Wildlife Art | Bfoxyarts"
    description = settings.get("Meta Description") or "See current, upcoming, and past exhibitions featuring original artwork by Bernadette Fox."
    canonical = settings.get("Canonical URL") or "https://bfoxyarts.com/exhibitions"
    email = settings.get("Contact Email") or "hello@bfoxyarts.com"
    source_cards = "\n".join(cards)
    return f'''<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>{escape(title)}</title>
<meta name="description" content="{escape(description, quote=True)}">
<link rel="canonical" href="{escape(canonical, quote=True)}">
<meta name="robots" content="index, follow, max-image-preview:large">
<meta name="author" content="Bernadette Fox, Bfoxyarts">
<meta property="og:type" content="website">
<meta property="og:title" content="{escape(title, quote=True)}">
<meta property="og:description" content="{escape(description, quote=True)}">
<meta property="og:url" content="{escape(canonical, quote=True)}">
<meta property="og:image" content="https://bfoxyarts.com/images/MFA2.jpg">
<link rel="icon" href="favicon.ico" sizes="any">
<link rel="icon" href="favicon-32x32.png" sizes="32x32" type="image/png">
<link rel="icon" href="favicon-16x16.png" sizes="16x16" type="image/png">
<link rel="apple-touch-icon" href="apple-touch-icon.png">
<link rel="manifest" href="site.webmanifest">
<link rel="stylesheet" href="style.css">
<style>
.exhibitions-shell{{padding:2rem 0 4rem}}.exhibition-tabs{{display:flex;justify-content:center;gap:.65rem;flex-wrap:wrap;margin:0 auto 2.2rem;padding:.45rem;max-width:690px;background:#f4eff7;border:1px solid rgba(91,44,131,.14);border-radius:999px}}.exhibition-tab{{appearance:none;border:0;background:transparent;color:#4c3d55;padding:.8rem 1.25rem;border-radius:999px;font:inherit;font-weight:700;cursor:pointer}}.exhibition-tab[aria-selected="true"]{{background:#5b2c83;color:#fff;box-shadow:0 5px 16px rgba(91,44,131,.25)}}.tab-count{{display:inline-grid;place-items:center;min-width:1.5rem;height:1.5rem;margin-left:.35rem;padding:0 .35rem;border-radius:999px;background:rgba(255,255,255,.22);font-size:.78rem}}.exhibition-panel[hidden]{{display:none}}.section-heading{{text-align:center;margin:0 0 1.5rem}}.section-heading p{{max-width:720px;margin:.35rem auto 0;color:#655b69}}.exhibition-list{{display:grid;gap:2rem}}.exhibition-card{{display:grid;grid-template-columns:minmax(260px,.85fr) minmax(0,1.15fr);overflow:hidden;background:#fff;border:1px solid rgba(91,44,131,.14);border-radius:20px;box-shadow:0 12px 34px rgba(40,28,46,.09)}}.exhibition-card.featured{{border-color:rgba(91,44,131,.32);box-shadow:0 18px 45px rgba(91,44,131,.14)}}.exhibition-media,.exhibition-image-placeholder{{background:#eeeaf1;min-height:340px;display:grid;place-items:center;overflow:hidden}}.exhibition-media img{{width:100%;height:100%;max-height:520px;object-fit:cover;display:block}}.exhibition-media.contain img{{object-fit:contain;background:#f8f6f9}}.exhibition-media.duo{{grid-template-columns:1fr 1fr;gap:2px}}.exhibition-image-placeholder{{color:#5b2c83;font-weight:800;font-size:1.4rem}}.exhibition-content{{padding:clamp(1.4rem,3vw,2.4rem);display:flex;flex-direction:column;justify-content:center}}.status-row{{display:flex;gap:.55rem;align-items:center;flex-wrap:wrap;margin-bottom:.7rem}}.status-badge,.featured-badge{{display:inline-flex;width:max-content;border-radius:999px;padding:.38rem .7rem;font-size:.78rem;font-weight:800;letter-spacing:.06em;text-transform:uppercase}}.status-upcoming{{background:#eee4f7;color:#5b2c83}}.status-current{{background:#e2f3e8;color:#24623a}}.status-past{{background:#ecebed;color:#5c5660}}.featured-badge{{background:#faefd2;color:#785b16}}.exhibition-content h2{{margin:.1rem 0 .7rem;font-size:clamp(1.55rem,3vw,2.25rem);line-height:1.12}}.exhibition-meta{{display:grid;gap:.45rem;margin:0 0 1rem;color:#4f4654}}.exhibition-meta p{{margin:0}}.exhibition-description{{line-height:1.72;margin:0 0 1rem}}.artwork-note{{margin:.25rem 0 1rem;padding:.8rem 1rem;border-left:4px solid #5b2c83;background:#f8f4fa;border-radius:0 10px 10px 0}}.exhibition-actions{{display:flex;gap:.7rem;flex-wrap:wrap;margin-top:.35rem}}.text-link{{display:inline-flex;align-items:center;padding:.72rem .2rem;color:#5b2c83;font-weight:700;text-decoration:underline;text-underline-offset:3px}}.empty-state{{max-width:760px;margin:0 auto;text-align:center;padding:2.8rem 1.5rem;background:#f8f5fa;border:1px solid rgba(91,44,131,.13);border-radius:18px}}.exhibitions-cta{{margin-top:3rem;padding:2.2rem;text-align:center;background:#f4eff7;border-radius:20px}}.exhibitions-cta .exhibition-actions{{justify-content:center}}@media(max-width:800px){{.exhibition-card{{grid-template-columns:1fr}}.exhibition-media,.exhibition-image-placeholder{{min-height:250px;max-height:430px}}.exhibition-tabs{{border-radius:18px}}.exhibition-tab{{flex:1 1 130px}}}}
</style>
</head>
<body>
<header class="site-header"><div class="header-inner"><a class="brand" href="/"><img src="images/logo.png" alt="Bfoxyarts logo"><span>Bfoxyarts</span></a><nav class="site-nav" aria-label="Primary navigation"><a href="/">Home</a><a href="gallery">Gallery</a><a href="exhibitions" aria-current="page">Exhibitions</a><a href="about">About</a><a href="contact">Contact</a></nav></div></header>
<main><section class="page-hero"><img class="page-logo" src="images/logo.png" alt="Bfoxyarts logo"><span class="eyebrow">See the artwork in person</span><h1>Exhibitions &amp; Gallery Appearances</h1><p class="lead" style="margin:0 auto">Explore current, upcoming, and past exhibitions featuring original artwork by Bernadette Fox.</p></section>
<section class="exhibitions-shell"><div class="container">
<nav class="exhibition-tabs" aria-label="Exhibition status" role="tablist"><button class="exhibition-tab" id="tab-upcoming" data-panel="upcoming" aria-controls="panel-upcoming" aria-selected="true" role="tab" type="button">Upcoming <span class="tab-count">0</span></button><button class="exhibition-tab" id="tab-current" data-panel="current" aria-controls="panel-current" aria-selected="false" role="tab" type="button">Current <span class="tab-count">0</span></button><button class="exhibition-tab" id="tab-past" data-panel="past" aria-controls="panel-past" aria-selected="false" role="tab" type="button">Past <span class="tab-count">0</span></button></nav>
<section class="exhibition-panel" id="panel-upcoming" aria-labelledby="tab-upcoming" role="tabpanel"><div class="section-heading"><h2>Upcoming Exhibitions</h2><p>Plan a visit to experience Bernadette's artwork in person.</p></div><div class="exhibition-list"></div><div class="empty-state" hidden><h2>No upcoming exhibitions are currently listed</h2></div></section>
<section class="exhibition-panel" id="panel-current" aria-labelledby="tab-current" role="tabpanel" hidden><div class="section-heading"><h2>Current Exhibitions</h2></div><div class="exhibition-list"></div><div class="empty-state" hidden><h2>No exhibitions are currently open</h2></div></section>
<section class="exhibition-panel" id="panel-past" aria-labelledby="tab-past" role="tabpanel" hidden><div class="section-heading"><h2>Past Exhibitions</h2><p>A record of gallery appearances and group exhibitions featuring Bernadette's artwork.</p></div><div class="exhibition-list"></div><div class="empty-state" hidden><h2>No past exhibitions are currently listed</h2></div></section>
<div id="exhibition-source" hidden>{source_cards}</div>
<section class="exhibitions-cta"><span class="eyebrow">Original wildlife art</span><h2>Interested in an artwork or commission?</h2><p>Browse available paintings or contact Bernadette Fox to discuss a purchase or commissioned piece.</p><div class="exhibition-actions"><a class="button" href="gallery">View Artwork Gallery</a><a class="text-link" href="contact">Contact the Artist</a></div></section>
</div></section></main>
<footer class="site-footer"><div class="footer-inner"><p class="footer-title">Interested in purchasing artwork or commissioning a custom piece?</p><p><a href="mailto:{escape(email, quote=True)}">Email Bernadette Fox</a></p><p><a href="https://www.instagram.com/bfoxyarts/" target="_blank" rel="me noopener noreferrer">Follow @bfoxyarts on Instagram</a></p><p class="footer-small">© {date.today().year} Bfoxyarts · Wildlife Art by Bernadette Fox</p></div></footer>
<script>
(function(){{const names=['upcoming','current','past'];const tabs=Array.from(document.querySelectorAll('.exhibition-tab'));const panels=Object.fromEntries(names.map(n=>[n,document.getElementById('panel-'+n)]));const today=new Date();today.setHours(0,0,0,0);function localDate(v){{const p=v.split('-').map(Number);return new Date(p[0],p[1]-1,p[2]);}}function statusFor(card){{const start=localDate(card.dataset.start),end=localDate(card.dataset.end);if(today<start)return'upcoming';if(today>end)return'past';return'current';}}function openPanel(name,focus){{tabs.forEach(tab=>{{const active=tab.dataset.panel===name;tab.setAttribute('aria-selected',String(active));tab.tabIndex=active?0:-1;if(active&&focus)tab.focus();}});names.forEach(key=>panels[key].hidden=key!==name);}}const cards=Array.from(document.querySelectorAll('#exhibition-source .exhibition-card'));cards.forEach(card=>{{const status=statusFor(card),badge=card.querySelector('.status-badge');badge.textContent=status[0].toUpperCase()+status.slice(1);badge.classList.add('status-'+status);panels[status].querySelector('.exhibition-list').appendChild(card);}});names.forEach(name=>{{const list=panels[name].querySelector('.exhibition-list');Array.from(list.children).sort((a,b)=>name==='past'?localDate(b.dataset.start)-localDate(a.dataset.start):localDate(a.dataset.start)-localDate(b.dataset.start)).forEach(card=>list.appendChild(card));const count=list.children.length;document.querySelector('[data-panel="'+name+'"] .tab-count').textContent=count;panels[name].querySelector('.empty-state').hidden=count!==0;}});document.getElementById('exhibition-source').remove();tabs.forEach((tab,index)=>{{tab.addEventListener('click',()=>openPanel(tab.dataset.panel,false));tab.addEventListener('keydown',event=>{{if(!['ArrowRight','ArrowLeft'].includes(event.key))return;event.preventDefault();const next=tabs[(index+(event.key==='ArrowRight'?1:-1)+tabs.length)%tabs.length];openPanel(next.dataset.panel,true);}});}});const initial=panels.current.querySelector('.exhibition-card')?'current':(panels.upcoming.querySelector('.exhibition-card')?'upcoming':'past');openPanel(initial,false);}}());
</script>
</body></html>'''


def main() -> int:
    if not WORKBOOK.exists():
        print(f"Workbook not found: {WORKBOOK}", file=sys.stderr)
        return 1
    try:
        workbook = load_workbook(WORKBOOK, data_only=False, read_only=True)
        cards = load_cards(workbook)
        settings = load_settings(workbook)
        html = build_html(cards, settings)
        OUTPUT.write_text(html, encoding="utf-8", newline="\n")
        print(f"Generated {OUTPUT.name} with {len(cards)} published exhibitions.")
        return 0
    except Exception as exc:
        print(f"Generation failed: {exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
